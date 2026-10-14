import re

from pptx.oxml.ns import qn

from engine.core import Check, Finding, Status
from engine.core import content, rules
from engine.core.shapes import PLACEHOLDER


class PlaceholderSlide(Check):
    id = "C002"
    title = "Placeholder or empty slide"
    group = "content"
    description = "Slides whose only text is 'Video here' / 'Insert video', or slides with nothing on them."
    report_only = True
    sources = ["any"]

    def detect(self, deck, ctx):
        phrases = {_norm(p) for p in rules.load("template_text")["slide_placeholders"]}
        by_slide: dict[int, list] = {}
        for sv in deck.iter_shapes():
            if sv.group_depth == 0:
                by_slide.setdefault(sv.slide_number, []).append(sv)
        findings = []
        for n, slide in enumerate(deck.slides, 1):
            shapes = by_slide.get(n, [])
            visible = [sv for sv in shapes if not (sv.kind == PLACEHOLDER and not sv.text_bearing)]
            text = _norm(" ".join(content.shape_text(sv.shape) for sv in shapes))
            if not visible and not _has_backdrop(slide):
                findings.append(Finding(self.id, Status.YELLOW,
                    "Slide has nothing on it. Fine if it's a deliberate black/walk-in slide.",
                    slide=n, subject="empty"))
            elif visible and not [sv for sv in visible if not (sv.text_bearing and _invisible(sv, slide))]:
                # Everything on it is text in the background's own colour (adversarial round 1: white on white).
                findings.append(Finding(self.id, Status.YELLOW,
                    "Slide looks empty in the show: its only text is the same colour as the background.",
                    slide=n, subject="invisible", evidence={"rule": "every text run's colour equals the slide's "
                                                                    "solid background colour"}))
            elif text in phrases:
                pictures = [sv for sv in visible if not sv.text_bearing]
                findings.append(Finding(self.id, Status.RED if not pictures else Status.YELLOW,
                    f"Slide only says '{text}'. The real content was probably never added.",
                    slide=n, subject="placeholder", evidence={"text": text, "other_shapes": len(pictures)}))
        return findings


def _solid_bg(slide) -> str | None:
    bg = slide._element.cSld.find(qn("p:bg"))
    clr = bg.find(f"{qn('p:bgPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}") if bg is not None else None
    return clr.get("val", "").upper() if clr is not None else None


def _invisible(sv, slide) -> bool:
    """Every run with text has an explicit colour equal to the slide's explicit solid background. Theme colours
    and pictures aren't resolved: unknown counts as visible."""
    bg = _solid_bg(slide)
    if not bg:
        return False
    runs = [r for r in sv.shape._element.iter(qn("a:r")) if (r.findtext(qn("a:t")) or "").strip()]
    if not runs:
        return False
    for r in runs:
        clr = r.find(f"{qn('a:rPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
        if clr is None or clr.get("val", "").upper() != bg:
            return False
    return True


def _has_backdrop(slide) -> bool:
    """Walk-in/out slides are often 'empty' but show a background image or layout artwork."""
    if slide._element.cSld.find(qn("p:bg")) is not None:
        return True
    for source in (slide.slide_layout, slide.slide_layout.slide_master):
        if source._element.cSld.find(qn("p:bg")) is not None and source is slide.slide_layout:
            return True
        if any(sh._element.tag == qn("p:pic") for sh in source.shapes):
            return True
    return False


def _norm(text: str) -> str:
    return re.sub(r"[^\w ]+", "", " ".join(text.lower().split())).strip()
