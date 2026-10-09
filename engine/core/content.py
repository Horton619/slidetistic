"""Shared read-only helpers for content, link and media checks."""
from __future__ import annotations

import re
from dataclasses import dataclass

from pptx.oxml.ns import qn

from . import rules

SHAPE_TAGS = {qn(t) for t in ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp")}
URL_RE = re.compile(r"""(?:https?://|www\.)[^\s<>"'”’)\]]+""", re.I)


def shape_text(shape) -> str:
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join(c.text for row in shape.table.rows for c in row.cells)
    return shape.text_frame.text if getattr(shape, "has_text_frame", False) and shape.has_text_frame else ""


def notes_text(slide) -> str:
    # slide.notes_slide would create a notes page on a slide without one; never call it blind.
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""


def is_hidden(slide) -> bool:
    return slide._element.get("show") == "0"


def slide_name(slide, number: int) -> str:
    """Name for humans (SPEC §7.4): cSld name, else title, else first text, else 'Slide N'."""
    name = slide._element.cSld.get("name")
    if name:
        return name
    title = slide.shapes.title
    if title is not None and title.has_text_frame and title.text_frame.text.strip():
        return _trim(title.text_frame.text)
    for shape in slide.shapes:
        text = shape_text(shape).strip()
        if text:
            return _trim(text)
    return f"Slide {number}"


def _trim(text: str, n: int = 60) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1] + "…"


def owning_shape(el) -> tuple[int | None, str | None]:
    """(shape id, shape name) of the slide shape that contains XML element `el`."""
    node = el
    while node is not None and node.tag not in SHAPE_TAGS:
        node = node.getparent()
    if node is None:
        return None, None
    c_nv = next(node.iter(qn("p:cNvPr")), None)
    return (int(c_nv.get("id")), c_nv.get("name")) if c_nv is not None else (None, None)


@dataclass
class Link:
    slide: int
    shape_id: int | None
    shape_name: str | None
    url: str
    where: str          # "slide text" | "notes" | "hyperlink" | "picture link" | "embedded"


def text_urls(deck) -> list[Link]:
    """URLs typed as text on slides and in speaker notes."""
    out = []
    for n, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            for url in URL_RE.findall(shape_text(shape)):
                out.append(Link(n, shape.shape_id, shape.name, _clean(url), "slide text"))
        for url in URL_RE.findall(notes_text(slide)):
            out.append(Link(n, None, None, _clean(url), "notes"))
    return out


def hyperlinks(deck) -> list[Link]:
    """Clickable external hyperlinks on slides (text and shape links)."""
    out = []
    for n, slide in enumerate(deck.slides, 1):
        part = slide.part
        for el in part._element.iter(qn("a:hlinkClick"), qn("a:hlinkHover")):
            r_id = el.get(qn("r:id"))
            if not r_id or r_id not in part.rels:
                continue
            rel = part.rels[r_id]
            if not rel.is_external:
                continue
            sid, sname = owning_shape(el)
            on_picture = el.getparent().tag == qn("p:cNvPr") and el.getparent().getparent().getparent().tag == qn("p:pic")
            out.append(Link(n, sid, sname, rel.target_ref, "picture link" if on_picture else "hyperlink"))
    return out


@dataclass
class ExternalRel:
    part: object
    r_id: str
    reltype: str
    target: str
    slide: int | None


def external_rels(deck) -> list[ExternalRel]:
    slide_no = {id(s.part): i + 1 for i, s in enumerate(deck.slides)}
    out = []
    for part in deck.prs.part.package.iter_parts():
        for r_id, rel in part.rels.items():
            if rel.is_external:
                out.append(ExternalRel(part, r_id, rel.reltype, rel.target_ref, slide_no.get(id(part))))
    return out


def media_kind(url: str) -> str | None:
    """Label from rules/media_links.yaml ('YouTube', 'Google Drive', …) or None."""
    for entry in rules.load("media_links")["sites"]:
        if re.search(entry["pattern"], url, re.I):
            return entry["label"]
    if re.search(rules.load("media_links")["file_extensions"], url.split("?")[0], re.I):
        return "Media file"
    return None


def is_video_site(url: str) -> bool:
    return any(re.search(e["pattern"], url, re.I) for e in rules.load("media_links")["sites"] if e.get("video"))


def is_local_path(target: str) -> bool:
    return any(re.search(p, target, re.I) for p in rules.load("media_links")["local_paths"])


def _clean(url: str) -> str:
    return url.rstrip(".,;:!?")
