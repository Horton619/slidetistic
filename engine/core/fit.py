"""Fit fixes, in David's order (CLAUDE.md rule 6): move/resize the box, then tighten line spacing,
then shrink the font as a last resort. All sizes come from PowerPoint's own render (ctx.layout); this
module only does geometry and edits.
"""
from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, field

from pptx.oxml.ns import qn

from .text_geometry import alignments, insets

RUN_PROPS = (qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr"))


def safe_rect(deck, pct: float) -> tuple[float, float, float, float]:
    W, H = deck.slide_width, deck.slide_height
    return W * pct / 100, H * pct / 100, W * (1 - pct / 100), H * (1 - pct / 100)


def place(edits, sv, *, left=None, top=None, width=None, height=None):
    """Move/resize using slide coordinates; converts into the shape's own (group) coordinates."""
    xf = sv.xf
    kw = {}
    if left is not None:
        kw["left"] = (left - xf.off_x) / xf.scale_x
    if top is not None:
        kw["top"] = (top - xf.off_y) / xf.scale_y
    if width is not None:
        kw["width"] = width / xf.scale_x
    if height is not None:
        kw["height"] = height / xf.scale_y
    edits.set_geometry(sv, **kw)


def grow(sv, extra: float, safe, fixed: "_Fixed | None" = None) -> tuple[float, float] | None:
    """New (left, width) after widening (or, negative, narrowing) by `extra` the way the text is
    aligned (SPEC §7.2): centred → both sides, left → to the right, right → to the left.
    No clamping of the BOX: a box's edge isn't where its text is (PAT 18: a box 0.8 in off the slide
    with 1.6 in of indent got shoved 1.8 in right). Callers check the predicted TEXT instead."""
    f = fixed or _Fixed(sv)
    b, ins, algn = f.bounds, f.ins, f.algn
    width = b.width + extra
    sl, _, sr, _ = safe
    if width <= 0 or (extra > 0 and width - ins["l"] - ins["r"] > sr - sl):
        return None
    if algn <= {"r"}:
        left = b.left - extra
    elif algn <= {"l", "just"}:
        left = b.left
    else:
        left = b.left - extra / 2
    return left, width


@dataclass
class _Fixed:
    """A shape's box, insets, alignment and id: constant while a fix is planned (nothing is edited
    until the plan is chosen), but each costs an XML lookup, and plan_fit asks for them hundreds of
    thousands of times on a big deck (PAT 23: 842k alignment reads)."""
    bounds: object
    ins: dict
    algn: set
    sid: int

    def __init__(self, sv):
        self.bounds, self.ins = sv.bounds, insets(sv.shape)
        self.algn, self.sid = alignments(sv.shape) or {"l"}, sv.shape_id


def text_area(sv) -> float:
    ins = insets(sv.shape)
    return sv.bounds.width - ins["l"] - ins["r"]


def word_gap(line) -> float:
    """Typical space between words on a rendered line (falls back to a third of the line height)."""
    gaps = [b.left - a.right for a, b in zip(line.boxes, line.boxes[1:]) if b.left > a.right]
    return sorted(gaps)[len(gaps) // 2] if gaps else line.height / 3


def _tx_body(sv):
    return sv.shape.text_frame._txBody


def scale_spacing(edits, sv, factor: float, label: str):
    """Tighten line spacing to `factor` of what it is (exact points or percent; adds a percent if unset)."""
    old = _tx_body(sv)
    new = deepcopy(old)
    for p in new.iter(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        if ppr is None:
            ppr = p.makeelement(qn("a:pPr"), {})
            p.insert(0, ppr)
        ln = ppr.find(qn("a:lnSpc"))
        if ln is None:
            ln = ppr.makeelement(qn("a:lnSpc"), {})
            ln.append(ln.makeelement(qn("a:spcPct"), {"val": "100000"}))
            ppr.insert(0, ln)
        for v in ln:
            v.set("val", str(int(int(v.get("val")) * factor)))
    edits.replace_element(sv.shape.part, old, new, prop="line spacing",
                          before="100%", after=f"{factor:.0%} ({label})")


def scale_font(edits, sv, factor: float, label: str) -> bool:
    """Shrink every explicit font size (and exact line spacing with it). False when sizes aren't explicit."""
    old = _tx_body(sv)
    new = deepcopy(old)
    sized = [el for el in new.iter(*RUN_PROPS) if el.get("sz")]
    if not sized:
        return False
    for el in sized:
        el.set("sz", str(max(100, int(int(el.get("sz")) * factor))))
    for pts in new.iter(qn("a:spcPts")):
        if pts.getparent().tag == qn("a:lnSpc"):
            pts.set("val", str(int(int(pts.get("val")) * factor)))
    edits.replace_element(sv.shape.part, old, new, prop="font size",
                          before="100%", after=f"{factor:.0%} ({label})")
    return True


# --- collisions: every move/widen is checked against the other text on the slide first ------------

def letter_zone(line) -> tuple[float, float]:
    """Top and bottom of where a line's letters are: from cap height to the baseline. A line box also holds
    the font's empty ascent above the capitals and the descender space below (about 17 % and 19 % of it for
    common fonts), so stacked labels' boxes overlap by a third with no letters meeting (David 2026-10-12:
    "IL" over "331" on a map). Descenders reaching into the next line's empty space don't count."""
    from . import rules
    th = rules.load("thresholds")
    h = line.bottom - line.top
    return line.top + h * th.get("letter_top_pct", 17) / 100, line.top + h * th.get("letter_baseline_pct", 81) / 100


def _word_spans(line, tol) -> list[tuple[float, float]]:
    """The line's words left→right when they're known and sit inside the line, else the whole line. Two text
    boxes often share one line in Canva/Google exports, their words interleaving ("First 10 | Elementary |
    School Hub"): the line boxes overlap, the words never do (David 2026-10-12)."""
    if line.boxes and line.boxes[0].left >= line.left - tol and line.boxes[-1].right <= line.right + tol:
        return [(w.left, w.right) for w in line.boxes]
    return [(line.left, line.right)]


def lines_collide(la, lb, tol) -> bool:
    """Same rule as L010 (without a picture): a word of one overlaps a word of the other sideways, and the
    letter zones overlap vertically."""
    if min(la.right, lb.right) - max(la.left, lb.left) <= tol:
        return False
    (ta, ba), (tb, bb) = letter_zone(la), letter_zone(lb)
    if min(ba, bb) - max(ta, tb) <= tol:
        return False
    return any(min(ra, rb) - max(lft_a, lft_b) > tol
               for lft_a, ra in _word_spans(la, tol) for lft_b, rb in _word_spans(lb, tol))


def collision(slide_layout: dict, sid: int, predicted: list, tol) -> int | None:
    """The other shape whose drawn text the predicted lines would hit, or None."""
    for other, st in slide_layout.items():
        if other == sid:
            continue
        if any(lines_collide(p, q, tol) for p in predicted for q in st.lines):
            return other
    return None


def shifted(lines, dx: float = 0, dy: float = 0) -> list:
    from .layout.pdf import Line, Word
    return [Line(l.left + dx, l.top + dy, l.right + dx, l.bottom + dy, l.words,
                 [Word(w.text, w.left + dx, w.right + dx) for w in l.boxes]) for l in lines]


def merged_line(top_line, width: float, sv, new_left: float, new_width: float):
    """Where a line `width` wide lands in a box at (new_left, new_width), by the text's alignment."""
    from .layout.pdf import Line
    ins = insets(sv.shape)
    algn = alignments(sv.shape) or {"l"}
    if algn <= {"r"}:
        right = new_left + new_width - ins["r"]
        left = right - width
    elif algn <= {"l", "just"}:
        left = new_left + ins["l"]
    else:
        left = new_left + ins["l"] + ((new_width - ins["l"] - ins["r"]) - width) / 2
    return Line(left, top_line.top, left + width, top_line.bottom, top_line.words, top_line.boxes)



# --- re-wrap prediction: where lines will break at another box width ----------------------------
# Uses PowerPoint's own word widths from the render (not our font maths, CLAUDE.md rule 4) and is
# calibrated against the breaks PowerPoint actually made. A fresh render still verifies every fix.

class Reflow:
    """scale: font size factor (word widths and line pitch); spacing: line-spacing factor (pitch only)."""

    def __init__(self, sv, st, scale: float = 1.0, spacing: float = 1.0, band: float = 0.01):
        self.sv, self.st, self.scale, self.spacing = sv, st, scale, spacing
        order = list(dict.fromkeys(st.paragraph_of_line))
        self.paras = []
        for p in order:
            lines = [l for l, q in zip(st.lines, st.paragraph_of_line) if q == p]
            self.paras.append(lines)
        gaps = [b.left - a.right for l in st.lines for a, b in zip(l.boxes, l.boxes[1:]) if b.left > a.right]
        self.gap = sorted(gaps)[len(gaps) // 2] if gaps else (st.lines[0].height / 3 if st.lines else 0)
        heights = [l.height for l in st.lines]
        self.line_h = sorted(heights)[len(heights) // 2] if heights else 0
        pitches = [b.top - a.top for lines in self.paras for a, b in zip(lines, lines[1:]) if b.top > a.top]
        self.pitch = sorted(pitches)[len(pitches) // 2] if pitches else self.line_h * 1.15
        self.area0 = text_area(sv)
        self.bounds, self.algn = sv.bounds, alignments(sv.shape) or {"l"}
        self.offsets = self._calibrate(band)

    def _w(self, word) -> float:
        return (word.right - word.left) * self.scale

    def _calibrate(self, band: float) -> tuple[float, float]:
        """The range of usable text widths PowerPoint's actual breaks allow, as offsets from the box's
        text area. Every line fits (≥ lo) and no line could have taken the next line's first word
        (< hi). The true width is near the box's own text area, so the range is that area ± `band`,
        clamped into [lo, hi]. A fix must work at BOTH ends: a single "narrowest possible" guess made
        narrowing look like it re-wraps when PowerPoint still had slack, so L009 crept ~1% a round
        to the 6-round cap (Loop Lid, CLARB 11; optimization round 2026-10-11)."""
        lo, hi = 0.0, float("inf")
        for lines in self.paras:
            for k, line in enumerate(lines):
                lo = max(lo, line.width)
                if k + 1 < len(lines) and lines[k + 1].boxes:
                    hi = min(hi, line.width + self.gap + (lines[k + 1].boxes[0].right - lines[k + 1].boxes[0].left))
        lo_w = lo * 1.002 if lo else self.area0
        hi_w = max(hi * 0.998 if hi != float("inf") else max(lo_w, self.area0), lo_w)
        a = max(lo_w, min(self.area0 * (1 - band), hi_w))
        b = min(hi_w, max(self.area0 * (1 + band), lo_w))
        return a - self.area0, b - self.area0

    def wrap(self, words, area) -> list[list]:
        out, cur, used = [], [], 0.0
        for w in words:
            ww = self._w(w)
            if cur and used + self.gap * self.scale + ww > area:
                out.append(cur)
                cur, used = [], 0.0
            used = ww if not cur else used + self.gap * self.scale + ww
            cur.append(w)
        if cur:
            out.append(cur)
        return out

    def predict(self, new_left: float, new_width: float):
        """(predicted Line list, orphan count, total line count) for the box at (new_left, new_width),
        at both ends of the calibrated range: lines from both (so collision and edge checks cover
        either outcome), the worse orphan and line counts."""
        a = self._predict_at(new_left, new_width, self.offsets[0])
        if self.offsets[1] == self.offsets[0]:
            return a
        b = self._predict_at(new_left, new_width, self.offsets[1])
        return a[0] + b[0], max(a[1], b[1]), max(a[2], b[2])

    def _predict_at(self, new_left: float, new_width: float, offset: float):
        from .layout.pdf import Line
        b = self.bounds
        area = (new_width - (b.width - self.area0)) + offset
        algn = self.algn
        out, orphans, total = [], 0, 0
        y = self.paras[0][0].top if self.paras and self.paras[0] else b.top
        for i, lines in enumerate(self.paras):
            words = [w for l in lines for w in l.boxes]
            wrapped = self.wrap(words, area)
            total += len(wrapped)
            if len(wrapped) >= 2 and len([w for w in wrapped[-1] if any(c.isalnum() for c in w.text)]) <= 1:
                orphans += 1
            first_off = lines[0].left - b.left
            cont_off = (lines[1].left - b.left) if len(lines) > 1 else first_off
            for k, ws in enumerate(wrapped):
                width = sum(self._w(w) for w in ws) + self.gap * self.scale * (len(ws) - 1)
                if algn <= {"r"}:
                    right = new_left + new_width - (b.right - lines[0].right)
                    left = right - width
                elif algn <= {"l", "just"}:
                    left = new_left + (first_off if k == 0 else cont_off)
                else:
                    left = new_left + new_width / 2 - width / 2
                top = y + k * self.pitch * self.scale * self.spacing
                out.append(Line(left, top, left + width, top + self.line_h * self.scale, [w.text for w in ws], ws))
            nxt = self.paras[i + 1][0].top if i + 1 < len(self.paras) and self.paras[i + 1] else None
            gap_after = (nxt - (lines[0].top + len(lines) * self.pitch)) if nxt is not None else 0
            y += len(wrapped) * self.pitch * self.scale * self.spacing + max(0.0, gap_after)
        return out, orphans, total


def is_graphic(sv) -> bool:
    """Images: pictures, plus shapes FILLED with an image. Canva exports photos and charts as
    image-filled freeforms (PAT 5's chart and map). Plain coloured panels behind text are not obstacles:
    counting them flagged 45 harmless overlaps on PAT."""
    from pptx.oxml.ns import qn
    from .shapes import OTHER, PICTURE
    if sv.bounds is None:
        return False
    if sv.kind == PICTURE:
        return True
    if sv.kind == OTHER and not sv.text_bearing:
        sp_pr = sv.shape.element.find(qn("p:spPr"))
        return sp_pr is not None and sp_pr.find(qn("a:blipFill")) is not None
    return False


def obstacles(deck, slide_number: int, sid: int, current_lines) -> list[tuple]:
    """Things this text doesn't overlap now, which a fix must not slide it onto: pictures (PAT 18: words fully
    hidden behind an illustration; PAT 5: a bullet pushed under a chart) and, since 2026-10-12, filled shapes,
    outlines and lines (David: a pull "hit other elements"; a title moved onto its green rule). Ones it already
    overlaps (a title on a photo, a text panel's own background) are fine. Detection (is_graphic, L011) still
    counts pictures only."""
    from .shapes import CONNECTOR, GROUP
    out = []
    slide_area = deck.slide_width * deck.slide_height
    thin = 3 * 12700            # a rule line has no height: give it 3 pt so text can't slide onto it
    for sv in deck.slide_shapes(slide_number):
        if sv.shape_id == sid or sv.bounds is None or sv.kind == GROUP:
            continue
        # Pictures always; and (David 2026-10-12: a pull "hit other elements", a title moved onto its green rule)
        # filled shapes, outlines and lines, except backgrounds. Ones the text already sits on are fine.
        shape_like = sv.kind == CONNECTOR or has_fill(sv)
        if not is_graphic(sv) and not (shape_like and sv.bounds.width * sv.bounds.height < 0.5 * slide_area):
            continue
        b = sv.bounds
        box = (b.left, b.top, max(b.right, b.left + thin), max(b.bottom, b.top + thin))
        if b.right - b.left < thin:
            box = (b.left - thin / 2, box[1], b.left + thin / 2, box[3])
        if b.bottom - b.top < thin:
            box = (box[0], b.top - thin / 2, box[2], b.top + thin / 2)
        if not any(hits_box(l, box) for l in current_lines):
            out.append(box)
    return out


def hits_box(line, box, tol: float = 0) -> bool:
    return (min(line.right, box[2]) - max(line.left, box[0]) > tol and
            min(line.bottom, box[3]) - max(line.top, box[1]) > tol)


@dataclass
class Plan:
    left: float
    top: float
    width: float
    font: float = 1.0
    spacing: float = 1.0
    lines: list = field(default_factory=list)
    total: int = 0
    orphans: int = 0


def plan_fit(deck, sv, st, slide_layout, safe, tol, *, accept, narrow_pct=15, widen_pct=0,
             fonts=(1.0,), spacings=(1.0,), move=False, obstacle_boxes=(), band=0.01) -> "Plan | None":
    """The cheapest change, in David's order (box → line spacing → font), whose PREDICTED text: never
    leaves the slide, hits no other text, slides onto no new picture, passes `accept(orphans, lines)`,
    and is no further into the safe border than it is now. With `move` (the safe-zone fix) it must end
    up fully inside the safe zone. (Requiring that of every fix blocked PAT 5's one-line title, which
    already sat a little into the top border.)"""
    if not getattr(sv.shape, "has_text_frame", False):
        # Tables and other graphic frames keep their text in cells, and their frame width doesn't size the
        # columns: only moving the whole frame is a real option (PC suite 2026-10-11: L008 picked a smaller
        # font for a table on Loop Lid slide 6 and crashed on "GraphicFrame has no text_frame").
        fonts, spacings, narrow_pct, widen_pct = (1.0,), (1.0,), 0, 0
    sl, stop, sr, sb = safe
    W, H = deck.slide_width, deck.slide_height
    if not move and st.bounds:
        cl, ct, cr, cb = st.bounds
        sl, stop, sr, sb = min(sl, max(cl, 0)), min(stop, max(ct, 0)), max(sr, min(cr, W)), max(sb, min(cb, H))
    fixed = _Fixed(sv)
    old = fixed.bounds.width
    # Text that sits on a filled shape (a panel, a pill) stays inside it: a wider wrap that spills past the
    # shape is worse than the problem it fixes (David 2026-10-15, DCS slide 62: L009 widened 8% and the text left
    # its panel). Where the text already overhangs the shape, it may go no further than now.
    backing = backing_shape(deck, sv) if not move else None
    if backing is not None and st.lines:
        ob = backing.bounds
        keep_l = min(ob.left, min(x.left for x in st.lines)) - tol
        keep_r = max(ob.right, max(x.right for x in st.lines)) + tol
    steps = sorted(range(-narrow_pct, widen_pct + 1), key=lambda d: (abs(d), d < 0))
    for font in fonts:
        for spacing in spacings:
            r = Reflow(sv, st, font, spacing, band)
            for d in steps:
                spot = grow(sv, old * d / 100, safe, fixed)
                if not spot:
                    continue
                lines, orphans, total = r.predict(*spot)
                if not lines or not accept(orphans, total):
                    continue
                dx = dy = 0.0
                if move:
                    l, t = min(x.left for x in lines), min(x.top for x in lines)
                    rr, b = max(x.right for x in lines), max(x.bottom for x in lines)
                    if rr - l > sr - sl or b - t > sb - stop:
                        continue
                    pad = W * 0.002
                    dx = sl + pad - l if l < sl else (sr - pad - rr if rr > sr else 0.0)
                    dy = stop + pad - t if t < stop else (sb - pad - b if b > sb else 0.0)
                    lines = shifted(lines, dx, dy)
                if (min(x.left for x in lines) < sl or max(x.right for x in lines) > sr or
                        min(x.top for x in lines) < stop or max(x.bottom for x in lines) > sb):
                    continue
                if backing is not None and st.lines and (min(x.left for x in lines) < keep_l
                                                         or max(x.right for x in lines) > keep_r):
                    continue
                if collision(slide_layout, fixed.sid, lines, tol) is not None:
                    continue
                if any(hits_box(x, box, tol) for x in lines for box in obstacle_boxes):
                    continue
                return Plan(spot[0] + dx, fixed.bounds.top + dy, spot[1], font, spacing, lines, total, orphans)
    return None


def autofit(sv) -> bool:
    """The box is shrinking its text to fit (PowerPoint's "Shrink text on overflow", `a:normAutofit` with a
    scale). PowerPoint
    redoes that shrink whenever the box is edited or, on some versions, opened, so a wrap predicted from
    today's render isn't reliable after a width, font or spacing change (review 2026-10-13, 1.6)."""
    tf = getattr(sv.shape, "text_frame", None) if getattr(sv.shape, "has_text_frame", False) else None
    body = tf._txBody.find(qn("a:bodyPr")) if tf is not None else None
    auto = body.find(qn("a:normAutofit")) if body is not None else None
    # Only boxes PowerPoint is shrinking right now (it saved a scale): 182 suite boxes have autofit on, most
    # of them not shrinking; refusing all of those would drop fixes David approved.
    return auto is not None and (auto.get("fontScale") is not None or auto.get("lnSpcReduction") is not None)


def apply(edits, sv, plan: Plan, label: str, finding=None, ctx=None):
    from . import tuning
    from .edits import FixRefused
    # Off by default (David 2026-10-13): on CLARB slides the skipped fixes were good, and they held after an edit in
    # PowerPoint (which redoes the shrink). Kept as a switch in thresholds.yaml in case Windows behaves differently.
    skip = bool(ctx and ctx.thresholds.get("fit_skip_shrinking_autofit", False))
    if skip and autofit(sv) and (plan.font < 1 or plan.spacing < 1 or abs(plan.width - sv.bounds.width) > 1):
        raise FixRefused("this box shrinks its text to fit (autofit), so PowerPoint would re-wrap it after a "
                         "resize; left as it was")
    tuning.allow(edits, sv, tuning.describe_plan(edits.deck, sv, plan, finding, ctx, edits.check_id))
    if plan.font < 1 and not scale_font(edits, sv, plan.font, label):
        return False
    if plan.spacing < 1:
        scale_spacing(edits, sv, plan.spacing, label)
    place(edits, sv, left=plan.left, top=plan.top, width=plan.width)
    return True


def is_header(deck, sv, slide_layout) -> bool:
    """The slide's header: its largest explicitly-sized text, if short (≤ 12 words). Going from 1 line
    to 2 there breaks the designer's layout (David 2026-10-10); a lone text block (PAT 19) isn't one."""
    from .content import shape_text

    def biggest(shape) -> int:
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            return 0
        return max((int(el.get("sz")) for el in shape.text_frame._txBody.iter(*RUN_PROPS) if el.get("sz")), default=0)

    texts = [s for s in deck.slide_shapes(sv.slide_number) if s.text_bearing]
    if len([s for s in texts if s.shape_id in slide_layout and slide_layout[s.shape_id].lines]) < 2:
        return False
    top = max(texts, key=lambda s: biggest(s.shape), default=None)
    return top is not None and top.shape_id == sv.shape_id and len(shape_text(sv.shape).split()) <= 12


def header_limit(deck, sv, slide_layout, ctx) -> int | None:
    """Max lines a header may have: what it had in the ORIGINAL deck. None for non-headers."""
    if not is_header(deck, sv, slide_layout):
        return None
    return ctx.baseline_lines.get((sv.slide_number, sv.shape_id))


def choices(ctx, threshold) -> tuple[tuple, tuple]:
    """Font and line-spacing factors to try, cheapest first, within the caps in thresholds.yaml."""
    cap = 1 - threshold(ctx, "fit_max_font_shrink_pct") / 100
    floor = threshold(ctx, "fit_min_line_spacing_pct") / 100
    fonts = tuple(f for f in (1.0, 0.95, 0.9, 0.85, 0.8) if f >= cap - 1e-9)
    spacings = tuple(x for x in (1.0, 0.95, 0.9, 0.85) if x >= floor - 1e-9)
    return fonts, spacings


# --- where text sits (David 2026-10-12): on a shape, or in a header/footer band --------------------

FOOTER_PLACEHOLDERS = ("FOOTER", "SLIDE_NUMBER", "DATE", "HEADER")


def has_fill(sv) -> bool:
    """Visibly drawn: a picture, or a shape with a fill or an outline (own or from its theme style)."""
    from .shapes import PICTURE
    if sv.kind == PICTURE:
        return True
    el = sv.shape.element
    sp_pr = el.find(qn("p:spPr"))
    if sp_pr is not None:
        if any(sp_pr.find(qn(f"a:{t}")) is not None for t in ("solidFill", "gradFill", "blipFill", "pattFill")):
            return True
        ln = sp_pr.find(qn("a:ln"))
        if ln is not None and len(ln) and ln.find(qn("a:noFill")) is None:
            return True
        if sp_pr.find(qn("a:noFill")) is not None:
            return False
    style = el.find(qn("p:style"))
    ref = style.find(qn("a:fillRef")) if style is not None else None
    return ref is not None and ref.get("idx") not in (None, "0")


def backing_shape(deck, sv, min_overlap: float = 0.8):
    """The filled shape this text box sits on (≥ `min_overlap` of the box over it), smallest first; None.
    Backgrounds (half the slide or more) don't count. The text's own fill moves with it anyway.
    David 2026-10-12: moving text off its shape broke the design (votes: 1 👍, 58 👎)."""
    from .shapes import CONNECTOR, GROUP
    if sv.bounds is None:
        return None
    b = sv.bounds
    area = max(b.width * b.height, 1)
    slide_area = deck.slide_width * deck.slide_height
    best = None
    for o in deck.slide_shapes(sv.slide_number):
        if o.shape_id == sv.shape_id or o.bounds is None or o.kind in (GROUP, CONNECTOR) or o.rotation:
            continue
        ob = o.bounds
        o_area = ob.width * ob.height
        if o_area >= 0.5 * slide_area or not has_fill(o):
            continue
        w = min(b.right, ob.right) - max(b.left, ob.left)
        h = min(b.bottom, ob.bottom) - max(b.top, ob.top)
        if w > 0 and h > 0 and w * h / area >= min_overlap and (best is None or o_area < best[0]):
            best = (o_area, o)
    return best[1] if best else None


def edge_band(deck, sv, box, pct: float) -> str:
    """'top' / 'bottom' when the drawn text sits entirely in the header or footer band (`pct` of the slide
    height), or 'placeholder' for footer / slide-number / date / header placeholders; else ''.
    David 2026-10-12: header and footer text is placed on purpose; moving it in caused collisions."""
    if getattr(sv.shape, "is_placeholder", False) if sv is not None else False:
        try:
            if str(sv.shape.placeholder_format.type).split(" ")[0].split(".")[-1] in FOOTER_PLACEHOLDERS:
                return "placeholder"
        except (AttributeError, ValueError):
            pass
    if box is None:
        return ""
    H = deck.slide_height
    _, t, _, b = box
    if b <= H * pct / 100:
        return "top"
    if t >= H * (1 - pct / 100):
        return "bottom"
    return ""


def area_pct(deck, sv) -> float:
    """The text box's share of the slide, in percent."""
    if sv is None or sv.bounds is None:
        return 0.0
    return 100 * sv.bounds.width * sv.bounds.height / (deck.slide_width * deck.slide_height)
