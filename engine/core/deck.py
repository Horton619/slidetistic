from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.package import XmlPart
from pptx.oxml.ns import qn

from .geometry import Bounds, Transform
from .shapes import GROUP, classify, has_text

# Remember each slide's shape list between edits (David 2026-10-11, optimization round: PAT 23 layout
# repair 42 s → 16 s). "on" (default), "off", or "check": every cached answer is compared with a fresh
# walk and a mismatch stops the run (suite/test mode, proves nothing reads a stale list). Set by
# `--shape-cache` or SLIDETISTIC_SHAPE_CACHE. Any edit outside EditSession must call invalidate().
SHAPE_CACHE = os.environ.get("SLIDETISTIC_SHAPE_CACHE", "on")


class StaleShapeCache(AssertionError):
    pass


# Render copies (David 2026-10-11): PowerPoint draws a video or sound as its poster image, a separate
# picture part, so the copies SlideTistic hands PowerPoint carry a tiny valid stand-in of the same
# format instead of the real media. Renders are identical and a 500 MB deck saves and opens in a
# moment. The deck David gets back always keeps the real files.
MEDIA_STUBS = Path(__file__).resolve().parent.parent / "assets" / "media_stubs"
_STUB_EXT = {p.suffix.lower() for p in MEDIA_STUBS.glob("stub.*")}
_STORED_EXT = _STUB_EXT | {".jpg", ".jpeg", ".png", ".gif", ".mpeg", ".m4v"}   # already compressed


def stub_for(member: str) -> bytes | None:
    """The stand-in for a media file inside a deck, or None to keep it."""
    ext = Path(member).suffix.lower()
    if member.startswith("ppt/media/") and ext in _STUB_EXT:
        return (MEDIA_STUBS / f"stub{ext}").read_bytes()
    return None


def compression_for(member: str) -> int:
    import zipfile
    return zipfile.ZIP_STORED if Path(member).suffix.lower() in _STORED_EXT else zipfile.ZIP_DEFLATED


def replace_retrying(src, dst, wait: float = 10.0) -> None:
    """os.replace that waits out a brief lock. On Windows, Dropbox, OneDrive or Defender can hold a file that
    was just written, and renaming over it then fails with "Access is denied" (PC suite 2026-10-11: a repair
    failed at its last step writing review.json into Dropbox). Retries for up to `wait` seconds."""
    import time
    deadline, delay = time.monotonic() + wait, 0.1
    while True:
        try:
            os.replace(src, dst)
            return
        except PermissionError:
            if time.monotonic() > deadline:
                raise
            time.sleep(delay)
            delay = min(delay * 2, 1.0)


def write_render_copy(src: Path, dst: Path) -> None:
    """Copy a deck for PowerPoint to draw: media swapped for stand-ins, compressed media stored."""
    import zipfile
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", strict_timestamps=False) as zout:
        for info in zin.infolist():
            stub = stub_for(info.filename)
            zout.writestr(info.filename, stub if stub is not None else zin.read(info),
                          compress_type=compression_for(info.filename))


@dataclass
class ShapeView:
    slide_index: int               # 0-based
    shape: object                  # python-pptx shape
    kind: str
    bounds: Bounds | None          # absolute, EMU; None when position is unknown
    group_depth: int               # 0 = directly on the slide
    rotation: float
    flipped: bool
    text_bearing: bool
    xf: Transform = None           # maps the shape's own coordinates to slide coordinates

    @property
    def slide_number(self) -> int:
        return self.slide_index + 1

    @property
    def shape_id(self) -> int:
        return self.shape.shape_id

    @property
    def name(self) -> str:
        return self.shape.name


class Deck:
    def __init__(self, path: str | os.PathLike):
        self.path = Path(path)
        self.prs = Presentation(str(self.path))
        # python-pptx keeps some XML parts (theme, presProps) as raw bytes. We parse them on
        # demand and write them back in save(), so edits to them go through the same path.
        self._blob_roots: dict[str, tuple[object, etree._Element]] = {}
        self._shape_memo: dict[int, list[ShapeView]] = {}

    def invalidate(self):
        """The deck's XML changed: forget the remembered shape lists."""
        self._shape_memo.clear()

    @property
    def slide_width(self) -> int:
        return self.prs.slide_width

    @property
    def slide_height(self) -> int:
        return self.prs.slide_height

    @property
    def slides(self):
        # prs.slides calls get_or_add_sldIdLst(), which writes an empty <p:sldIdLst> into a
        # slideless deck. Reading must never change the deck, so check for it first.
        if self.prs.part._element.sldIdLst is None:
            return []
        return list(self.prs.slides)

    def iter_shapes(self) -> list[ShapeView]:
        """Every shape on every slide, groups included, children of groups flattened in."""
        if SHAPE_CACHE != "off":
            return [sv for n in range(1, len(self.slides) + 1) for sv in self.slide_shapes(n)]
        out: list[ShapeView] = []
        for i, slide in enumerate(self.slides):
            self._walk(i, slide.shapes, Transform(), 0, out)
        return out

    def _walk(self, slide_index, shapes, xf: Transform, depth: int, out: list):
        for shape in shapes:
            kind = classify(shape)
            bounds = None
            if None not in (shape.left, shape.top, shape.width, shape.height):
                bounds = xf.apply(shape.left, shape.top, shape.width, shape.height)
            xfrm = _xfrm(shape)
            out.append(ShapeView(
                slide_index=slide_index,
                shape=shape,
                kind=kind,
                bounds=bounds,
                group_depth=depth,
                rotation=getattr(shape, "rotation", 0.0) or 0.0,
                flipped=xfrm is not None and (xfrm.get("flipH") == "1" or xfrm.get("flipV") == "1"),
                text_bearing=has_text(shape),
                xf=xf,
            ))
            if kind == GROUP:
                child_xf = xf.child_transform(xfrm) if xfrm is not None else xf
                self._walk(slide_index, shape.shapes, child_xf, depth + 1, out)

    def slide_shapes(self, slide_number: int) -> list[ShapeView]:
        """iter_shapes() for one slide (1-based). Walking every slide to keep one made layout repair
        quadratic in slide count: each walk resolves placeholder inheritance through python-pptx."""
        slides = self.slides
        if not 1 <= slide_number <= len(slides):
            return []
        if SHAPE_CACHE != "off" and slide_number in self._shape_memo:
            hit = self._shape_memo[slide_number]
            if SHAPE_CACHE == "check":
                fresh: list[ShapeView] = []
                self._walk(slide_number - 1, slides[slide_number - 1].shapes, Transform(), 0, fresh)
                if [_facts(sv) for sv in fresh] != [_facts(sv) for sv in hit]:
                    raise StaleShapeCache(f"slide {slide_number}: remembered shapes differ from the deck")
            return list(hit)
        out: list[ShapeView] = []
        self._walk(slide_number - 1, slides[slide_number - 1].shapes, Transform(), 0, out)
        if SHAPE_CACHE != "off":
            self._shape_memo[slide_number] = out
            return list(out)
        return out

    def find_shape(self, slide_number: int, shape_id: int) -> ShapeView | None:
        for sv in self.slide_shapes(slide_number):
            if sv.shape_id == shape_id:
                return sv
        return None

    def xml_parts(self) -> list[tuple[object, etree._Element]]:
        """(part, root element) for every XML part, excluding embedded font data."""
        out = []
        for part in self.prs.part.package.iter_parts():
            name = str(part.partname)
            if not name.endswith(".xml") or name.startswith("/ppt/fonts/"):
                continue
            if isinstance(part, XmlPart):
                out.append((part, part._element))
            else:
                if name not in self._blob_roots:
                    self._blob_roots[name] = (part, etree.fromstring(part.blob))
                out.append(self._blob_roots[name])
        return out

    def slide_number_of(self, part) -> int | None:
        for i, slide in enumerate(self.slides):
            if slide.part is part:
                return i + 1
        return None

    def save(self, out_path: str | os.PathLike, render_copy: bool = False) -> Path:
        """Atomic: write to a temp name beside the target, then rename over it. `render_copy`: a copy
        only PowerPoint will draw (fit rounds): media written as stand-ins, see write_render_copy."""
        for part, root in self._blob_roots.values():
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        out_path = Path(out_path)
        tmp = out_path.with_name(out_path.name + ".partial")
        if render_copy:
            from pptx.opc import serialized
            original = serialized._ZipPkgWriter.write

            def write(writer, pack_uri, blob):
                name = pack_uri.membername
                stub = stub_for(name)
                writer._zipf.writestr(name, stub if stub is not None else blob, compress_type=compression_for(name))

            serialized._ZipPkgWriter.write = write
            try:
                self.prs.save(str(tmp))
            finally:
                serialized._ZipPkgWriter.write = original
        else:
            self.prs.save(str(tmp))
        replace_retrying(tmp, out_path)
        return out_path

    def part_blobs(self) -> dict[str, bytes]:
        """Serialized bytes of every part. Used to prove a step changed nothing."""
        return {str(p.partname): p.blob for p in self.prs.part.package.iter_parts()}


def _xfrm(shape):
    el = shape.element
    for path in ("p:spPr/a:xfrm", "p:grpSpPr/a:xfrm", "p:xfrm"):
        found = el.find("/".join(qn(t) for t in path.split("/")))
        if found is not None:
            return found
    return None


def _facts(sv: ShapeView) -> tuple:
    """What a check reads from a ShapeView, for the cache's check mode."""
    from lxml import etree as _et
    return (sv.shape_id, sv.kind, sv.bounds, sv.group_depth, sv.rotation, sv.flipped, sv.text_bearing,
            _et.tostring(sv.shape.element))

