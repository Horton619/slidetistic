"""Where text can be inside its box, from the box's inner margins, paragraph indents and anchoring.

No measuring (CLAUDE.md rule 4): this only uses values written in the file. Margins and indents are
text units and are NOT scaled by a group's transform. Values inherited from layouts/masters aren't
resolved; a missing local value counts as PowerPoint's default (insets) or 0 (indents).
"""
from __future__ import annotations

from pptx.oxml.ns import qn

DEFAULT_INSETS = {"l": 91440, "t": 45720, "r": 91440, "b": 45720}   # 0.1 / 0.05 in


def body_pr(shape):
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return None
    return shape.text_frame._txBody.find(qn("a:bodyPr"))


def insets(shape) -> dict[str, int]:
    bp = body_pr(shape)
    return {k: int(bp.get(f"{k}Ins", v)) if bp is not None else v for k, v in DEFAULT_INSETS.items()}


def anchor(shape) -> str:
    bp = body_pr(shape)
    return bp.get("anchor", "t") if bp is not None else "t"


def _paragraphs(shape):
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return []
    out = []
    for p in shape.text_frame._txBody.iter(qn("a:p")):
        if "".join(t.text or "" for t in p.iter(qn("a:t"))).strip():
            out.append(p.find(qn("a:pPr")))
    return out


def alignments(shape) -> set[str]:
    return {(ppr.get("algn", "l") if ppr is not None else "l") for ppr in _paragraphs(shape)}


def text_start_offset(shape) -> int:
    """Smallest distance from the box's left edge to where any line (or its bullet) starts."""
    ins = insets(shape)["l"]
    offsets = []
    for ppr in _paragraphs(shape):
        mar = int(ppr.get("marL", 0)) if ppr is not None else 0
        ind = int(ppr.get("indent", 0)) if ppr is not None else 0
        offsets.append(ins + mar + min(ind, 0))
    return min(offsets) if offsets else ins


def text_end_offset(shape) -> int:
    ins = insets(shape)["r"]
    offsets = [ins + (int(ppr.get("marR", 0)) if ppr is not None else 0) for ppr in _paragraphs(shape)]
    return min(offsets) if offsets else ins


def anchored_edge_off(shape, bounds, slide_w: int, tol: int) -> str | None:
    """'left'/'right' when text pinned to a box edge provably starts off the slide (words cut off)."""
    algn = alignments(shape)
    if algn and algn <= {"l", "just"} and bounds.left + text_start_offset(shape) < -tol:
        return "left"
    if algn and algn <= {"r"} and bounds.right - text_end_offset(shape) > slide_w + tol:
        return "right"
    return None


def text_overhang(shape, bounds, slide_w: int, slide_h: int) -> dict[str, int]:
    """How far the area text can occupy runs past each slide edge (EMU, >0 = past it).
    The top only counts for top-anchored text; the bottom only for bottom-anchored or full boxes is
    unknowable, so it's reported against the bottom inset alone."""
    ins = insets(shape)
    out = {"left": -(bounds.left + ins["l"]), "right": bounds.right - ins["r"] - slide_w,
           "bottom": bounds.bottom - ins["b"] - slide_h}
    out["top"] = -(bounds.top + ins["t"]) if anchor(shape) == "t" else 0
    return out
