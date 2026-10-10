"""Timing trees here are PowerPoint's own saves (tests/fixtures/media_modes, Mac + Windows)."""
import re

import pytest
from lxml import etree
from pptx.oxml.ns import qn

from engine.core import Deck
from engine.core.media import media_items
from tests.checks.helpers import Status, run_repair, run_scan
from tests.fixtures.builders import MEDIA_MODES, make_deck, text_box, video


@pytest.mark.parametrize("mode,expected", [
    ("mac_click_sequence", "click_sequence"), ("win_click_sequence", "click_sequence"),
    ("mac_automatic", "automatic"), ("win_automatic", "automatic"),
    ("mac_when_clicked", "on_click"), ("win_when_clicked", "on_click"),
    ("mac_toggle_in_sequence", "unknown"), ("mac_media_node_only", "no_timing")])
def test_reads_powerpoint_saved_play_modes(tmp_path, mode, expected):
    src = make_deck(tmp_path, video(), timing=mode)
    [m] = media_items(Deck(src))
    assert m.mode == expected


def test_m001_click_sequence_and_automatic_are_fine(tmp_path):
    # David 2026-10-09: clicker-driven shows want "In click sequence"; automatic is fine too.
    src = make_deck(tmp_path, slides=[[video()], [video()], [video()], [video()]],
                    timing={1: "win_click_sequence", 2: "mac_click_sequence", 3: "win_automatic", 4: "mac_automatic"})
    assert run_scan(src, "M001").findings == []


def test_m001_when_clicked_on_is_red_no_setting_yellow(tmp_path):
    src = make_deck(tmp_path, slides=[[video()], [video()]], timing={1: "win_when_clicked"})
    found = {f.slide: (f.status, f.evidence["mode"]) for f in run_scan(src, "M001").findings}
    assert found == {1: (Status.RED, "on_click"), 2: (Status.YELLOW, "no_timing")}


def _main_seq_shape(xml_root) -> str:
    """The main-sequence block with ids/durations stripped: what PowerPoint's structure looks like."""
    main = next(c for c in xml_root.iter(qn("p:cTn")) if c.get("nodeType") == "mainSeq").getparent()
    text = etree.tostring(main).decode()
    return re.sub(r' (id|dur|spid)="[^"]*"', "", re.sub(r"\s+xmlns:\w+=\"[^\"]+\"", "", text))


@pytest.mark.parametrize("platform", ["mac", "win"])
def test_m001_fix_matches_powerpoints_own_click_sequence(tmp_path, platform):
    src = make_deck(tmp_path, video(), timing=f"{platform}_when_clicked")
    result = run_repair(src, "M001")
    assert [f.status for f in result.findings if f.check_id == "M001"] == [Status.GREEN_FIXED]
    [m] = media_items(Deck(result.output_path))
    assert (m.mode, m.clicks) == ("click_sequence", 1)
    fixed = Deck(result.output_path).slides[0]._element.find(qn("p:timing"))
    saved = etree.fromstring((MEDIA_MODES / f"{platform}_click_sequence.xml").read_text().split("-->", 1)[1].strip())
    assert _main_seq_shape(fixed) == _main_seq_shape(saved)
    ids = [c.get("id") for c in fixed.iter(qn("p:cTn"))]
    assert len(ids) == len(set(ids)), "cTn ids must stay unique"


def test_m001_fix_adds_last_click_after_existing_animations(tmp_path):
    # Slide already has a click sequence (another video): ours becomes the next click.
    from tests.fixtures.builders import set_timing
    def post(prs):
        slide = prs.slides[0]
        set_timing(slide, "win_click_sequence")                    # first video: click 1
        first, second = [sh for sh in slide.shapes if sh._element.find(".//" + qn("a:videoFile")) is not None]
        timing = slide._element.find(qn("p:timing"))
        inter = etree.fromstring((MEDIA_MODES / "win_when_clicked.xml").read_text().split("-->", 1)[1].strip()
                                 .replace("{SPID}", str(second.shape_id)))
        trigger_seq = next(s for s in inter.iter(qn("p:seq")) if s.find(qn("p:cTn")).get("nodeType") == "interactiveSeq")
        for c in trigger_seq.iter(qn("p:cTn")):
            c.set("id", str(int(c.get("id")) + 100))
        timing.find(qn("p:tnLst")).find(qn("p:par")).find(qn("p:cTn")).find(qn("p:childTnLst")).append(trigger_seq)
    src = make_deck(tmp_path, video(), video(x=7), post=post)
    result = run_repair(src, "M001")
    modes = sorted((m.mode, m.clicks) for m in media_items(Deck(result.output_path)))
    assert modes == [("click_sequence", 1), ("click_sequence", 2)]


def test_m001_fix_refuses_no_setting_videos(tmp_path):
    result = run_repair(make_deck(tmp_path, video()), "M001")
    assert result.changes == [] and result.findings[0].evidence["fix_result"].startswith("refused")
