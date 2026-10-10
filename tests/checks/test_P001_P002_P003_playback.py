from lxml import etree
from pptx.oxml.ns import qn

from engine.core import Deck
from engine.core import playback
from tests.checks.helpers import Status, run_repair, run_scan
from tests.fixtures.builders import make_deck, text_box

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def transition(slide, alternate=True, **attrs):
    """The layout PowerPoint 2010+ saves (corpus): a p14 Choice and a plain Fallback copy."""
    a = " ".join(f'{k}="{v}"' for k, v in attrs.items())
    if alternate:
        xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}" xmlns:p14="{P14}">'
               f'<mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="2000" {a}/></mc:Choice>'
               f'<mc:Fallback><p:transition spd="slow" {a}/></mc:Fallback></mc:AlternateContent>')
    else:
        xml = f'<p:transition xmlns:p="{P}" spd="slow" {a}/>'
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(xml))


def show_pr(prs, xml_inner="", **attrs):
    part = next(p for p in prs.part.package.iter_parts() if str(p.partname) == "/ppt/presProps.xml")
    root = etree.fromstring(part.blob)
    old = root.find(qn("p:showPr"))
    if old is not None:
        root.remove(old)
    a = " ".join(f'{k}="{v}"' for k, v in attrs.items())
    root.insert(0, etree.fromstring(f'<p:showPr xmlns:p="{P}" {a}>{xml_inner}</p:showPr>'))
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def test_p001_walk_in_loop_message(tmp_path):
    def post(prs):
        for s in prs.slides:
            transition(s, advClick="0", advTm="15000")
    src = make_deck(tmp_path, slides=[[text_box("a")], [text_box("b")]], post=post)
    [f] = run_scan(src, "P001").findings
    assert f.status == Status.YELLOW and "walk-in loop" in f.message


def test_p001_leftover_timings_and_fix_removes_both_copies(tmp_path):
    def post(prs):
        transition(prs.slides[1], advTm="54975")
        transition(prs.slides[2], alternate=False, advTm="27984")
    src = make_deck(tmp_path, slides=[[text_box("a")], [text_box("b")], [text_box("c")]], post=post)
    [f] = run_scan(src, "P001").findings
    assert f.evidence["slides"] == [2, 3] and "rehearsal" in f.message
    result = run_repair(src, "P001")
    out = Deck(result.output_path)
    assert all(t.get("advTm") is None for s in out.slides for t in playback.transitions(s))
    assert len([t for t in playback.transitions(out.slides[1])]) == 2     # Choice + Fallback both kept


def test_p001_silent_when_use_timings_off(tmp_path):
    def post(prs):
        transition(prs.slides[0], advTm="5000")
        show_pr(prs, useTimings="0")
    src = make_deck(tmp_path, text_box(), post=post)
    assert run_scan(src, "P001").findings == []
    assert [f.subject for f in run_scan(src, "P003").findings] == ["timings-off"]


def test_p002_click_off_without_timer(tmp_path):
    def post(prs):
        transition(prs.slides[0], advClick="0")
        transition(prs.slides[1], advClick="0", advTm="15000")     # look-alike: walk-in, not stuck
    src = make_deck(tmp_path, slides=[[text_box("a")], [text_box("b")]], post=post)
    assert [f.slide for f in run_scan(src, "P002").findings] == [1]


def test_p003_kiosk_loop_and_range(tmp_path):
    src = make_deck(tmp_path, text_box(), post=lambda prs: show_pr(
        prs, '<p:kiosk/><p:sldRg st="2" end="5"/>', loop="1", showAnimation="0"))
    found = {f.subject: f.status for f in run_scan(src, "P003").findings}
    assert found == {"kiosk": Status.RED, "loop": Status.YELLOW, "no-animation": Status.RED, "range": Status.RED}


def test_p003_normal_settings_clean(tmp_path):
    src = make_deck(tmp_path, text_box(), post=lambda prs: show_pr(prs, showNarration="1"))
    assert run_scan(src, "P003").findings == []
