from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from engine.core import Deck, EMU_PER_INCH
from tests.checks.helpers import Status, run_repair, run_scan
from tests.fixtures.builders import group, make_deck, picture, text_box


def aligned(add, align):
    def wrapped(shapes):
        tb = add(shapes)
        tb.text_frame.paragraphs[0].alignment = align
        return tb
    return wrapped


def g001(result):
    return [(f.shape_name, f.status) for f in result.findings if f.check_id == "G001"]


def test_g001_left_aligned_text_off_left_edge_is_red_and_fixed(tmp_path):
    src = make_deck(tmp_path, text_box("Agenda", x=-0.56, name="Body"))
    assert g001(run_scan(src, "G001")) == [("Body", Status.RED)]
    result = run_repair(src, "G001")
    assert g001(result) == [("Body", Status.GREEN_FIXED)]
    sv = next(s for s in Deck(result.output_path).iter_shapes() if s.name == "Body")
    assert round(sv.bounds.left / EMU_PER_INCH, 2) == 0.2


def test_g001_centered_overhang_is_yellow_and_left_alone(tmp_path):
    # PAT MBD slide 22: centered box at x = -0.56 in. The hand-made MacFix left it too.
    src = make_deck(tmp_path, aligned(text_box("Research findings", x=-0.56, w=11.9, name="TextBox 6"), PP_ALIGN.CENTER),
                    size=(20, 11.25))
    result = run_repair(src, "G001")
    assert g001(result) == [("TextBox 6", Status.YELLOW)] and result.changes == []


def test_g001_right_aligned_off_right_edge_is_red(tmp_path):
    src = make_deck(tmp_path, aligned(text_box("Q&A", x=11, w=3, name="R"), PP_ALIGN.RIGHT))
    assert g001(run_scan(src, "G001")) == [("R", Status.RED)]


def test_g001_pictures_off_slide_never_flagged(tmp_path):
    src = make_deck(tmp_path, picture(x=-2, w=5), picture(x=12, w=4), text_box("fine", x=1))
    assert g001(run_scan(src, "G001")) == []


def test_g001_group_child_reported_not_moved(tmp_path):
    src = make_deck(tmp_path, group(text_box("in group", x=-0.5, name="Inner")))
    result = run_repair(src, "G001")
    assert result.changes == [] and g001(result)[0][0] == "Inner"


def test_g002_below_bottom_and_g003_rotated(tmp_path):
    src = make_deck(tmp_path, text_box("low", y=7, h=1, name="Low"), text_box("tilt", rotation=350, name="Tilt"))
    assert [f.shape_name for f in run_scan(src, "G002").findings] == ["Low"]
    assert [f.shape_name for f in run_scan(src, "G003").findings] == ["Tilt"]
    assert g001(run_scan(src, "G001")) == []


def test_g004_one_informational_finding_per_group(tmp_path):
    src = make_deck(tmp_path, group(text_box("a"), text_box("b"), picture()), group(picture()))
    [f] = run_scan(src, "G004").findings
    assert f.informational and len(f.evidence["text_shapes"]) == 2


def test_g001_indented_bullets_are_not_certain(tmp_path):
    # PAT slide 18: box at x = -0.81 in, but marL 1.62 in / indent -0.54 in put bullets and words on the slide.
    def indented(shapes):
        tb = text_box("Home Visiting: a trusted relationship", x=-0.81, w=8, name="Bullets")(shapes)
        ppr = tb.text_frame.paragraphs[0]._p.get_or_add_pPr()
        ppr.set("marL", "1478914")
        ppr.set("indent", "-492971")
        return tb
    src = make_deck(tmp_path, indented)
    assert g001(run_scan(src, "G001")) == [("Bullets", Status.YELLOW)]


def test_g001_overhang_within_inner_margin_ignored(tmp_path):
    # PAT: boxes 0.04-0.14 in past the top, centered vertically: no text can be there.
    def centered_top(shapes):
        tb = text_box("Cognitive Behavioral Therapy", y=-0.12, h=1.5, name="Title")(shapes)
        tb.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
        return tb
    src = make_deck(tmp_path, centered_top, text_box("tiny", x=-0.05, name="Tiny"))   # within 0.1 in left margin
    assert g001(run_scan(src, "G001")) == []
