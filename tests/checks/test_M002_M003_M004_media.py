from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from tests.checks.helpers import Status, run_scan
from tests.fixtures.builders import make_deck, text_box, video


def test_m002_inventory_reads_powerpoint_settings(tmp_path):
    src = make_deck(tmp_path, slides=[[text_box("Opener"), video()], [video()]],
                    timing={1: "win_automatic", 2: "mac_click_sequence"})
    inv = {f.slide: f for f in run_scan(src, "M002").findings}
    assert inv[1].informational and inv[1].evidence["slide_name"] == "Opener"
    assert (inv[1].evidence["mode"], inv[1].evidence["loop"], inv[1].evidence["volume"]) == ("automatic", True, 80)
    assert (inv[2].evidence["mode"], inv[2].evidence["full_screen"], inv[2].evidence["hidden_when_stopped"]) == \
        ("click_sequence", True, True)


def test_m002_rewind_after_playing(tmp_path):
    # broken_mac.pptx slide 05: the only difference from click sequence is fill="remove" on the media node.
    src = make_deck(tmp_path, slides=[[video()], [video()]], timing={1: "mac_rewind", 2: "mac_click_sequence"})
    inv = {f.slide: f for f in run_scan(src, "M002").findings}
    assert inv[1].evidence["rewind"] and "rewinds" in inv[1].message
    assert inv[1].evidence["mode"] == "click_sequence"
    assert not inv[2].evidence["rewind"]


def test_m002_muted_video_called_out(tmp_path):
    src = make_deck(tmp_path, video(), timing="mac_automatic")   # this PowerPoint save has vol="0"
    [f] = run_scan(src, "M002").findings
    assert "MUTED" in f.message and f.evidence["volume"] == 0


def test_m003_linked_video_missing_is_red_and_m004_wmv_yellow(tmp_path):
    def post(prs):
        pic = prs.slides[0].shapes[0]
        vf = pic._element.find(".//" + qn("a:videoFile"))
        r_id = prs.slides[0].part.relate_to("file:///Users/pat/Desktop/opener.mov", RT.VIDEO, is_external=True)
        vf.set(qn("r:link"), r_id)
        for ext in pic._element.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}media"):
            ext.getparent().getparent().remove(ext.getparent())   # drop the embedded copy: linked only
    src = make_deck(tmp_path, video(), post=post)
    assert [f.status for f in run_scan(src, "M003").findings] == [Status.RED]

    src2 = make_deck(tmp_path, video(name="clip.wmv"), name="wmv.pptx",
                     post=lambda prs: _rename_media(prs, "/ppt/media/media1.wmv"))
    assert [f.status for f in run_scan(src2, "M004").findings] == [Status.YELLOW]


def test_m003_m004_embedded_mp4_clean(tmp_path):
    src = make_deck(tmp_path, video())
    assert run_scan(src, "M003").findings == [] and run_scan(src, "M004").findings == []


def _rename_media(prs, partname):
    from pptx.opc.packuri import PackURI
    part = next(p for p in prs.part.package.iter_parts() if str(p.partname).startswith("/ppt/media/media"))
    part.partname = PackURI(partname)
